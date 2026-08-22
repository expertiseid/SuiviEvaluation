"""
Import combiné « planification complète » : un seul classeur Excel pour créer
d'un coup un cadre stratégique (avec sa structuration), un projet (avec sa
planification Objectif général > Objectif spécifique > Activité >
Sous-activité), ses indicateurs et ses bénéficiaires — en réutilisant telles
quelles les fonctions d'import de chaque app (apps.strategy, apps.projects,
apps.indicators, apps.beneficiaries). Même philosophie « best-effort » que
ces imports unitaires.

Seule la feuille « Projet » est obligatoire : cadre stratégique, indicateurs
et bénéficiaires sont chacun optionnels (feuille absente ou vide = section
ignorée).
"""
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def _prefixer(erreurs, etiquette):
    return [{"ligne": e["ligne"], "message": f"[{etiquette}] {e['message']}"} for e in erreurs]


def generer_modele_import_complet() -> bytes:
    from apps.beneficiaries.services.import_excel import COLONNES as COLONNES_BENEFICIAIRES
    from apps.indicators.import_excel import generer_feuille_indicateurs, generer_feuille_paliers_alerte
    from apps.projects.import_excel import CHAMPS_PROJET, generer_feuille_planification

    wb = Workbook()

    feuille_cadre = wb.active
    feuille_cadre.title = "Cadre stratégique"
    for cle, exemple in [
        ("Nom du cadre stratégique", "Plan stratégique 2026-2030"),
        ("Description", "Cadre de référence de l'organisation pour la période — facultatif."),
    ]:
        feuille_cadre.append([cle, exemple])
    for ligne in feuille_cadre.iter_rows(min_col=1, max_col=1):
        ligne[0].font = Font(bold=True)
    feuille_cadre.column_dimensions["A"].width = 32
    feuille_cadre.column_dimensions["B"].width = 55

    feuille_struct = wb.create_sheet("Structuration")
    feuille_struct.append(["Orientation stratégique", "Axe d'intervention"])
    for cellule in feuille_struct[1]:
        cellule.font = Font(bold=True)
    feuille_struct.append(["Amplifier la diffusion des pratiques agroécologiques", ""])
    feuille_struct.append(["", "Renforcement des capacités des producteurs et productrices"])

    feuille_projet = wb.create_sheet("Projet")
    for cle, exemple in CHAMPS_PROJET:
        if cle == "Cadre stratégique":
            continue
        feuille_projet.append([cle, exemple])
    for ligne in feuille_projet.iter_rows(min_col=1, max_col=1):
        ligne[0].font = Font(bold=True)
    feuille_projet.column_dimensions["A"].width = 42
    feuille_projet.column_dimensions["B"].width = 48

    generer_feuille_planification(wb)
    generer_feuille_indicateurs(wb, avec_colonne_projet=False)

    feuille_benef = wb.create_sheet("Bénéficiaires")
    feuille_benef.append(COLONNES_BENEFICIAIRES)
    for cellule in feuille_benef[1]:
        cellule.font = Font(bold=True)
    feuille_benef.append(
        ["Traoré", "Awa", "F", "1990-05-12", "70000001", "CNIB123456", "CNIB", "BF", "Kadiogo", "Kadiogo", "Koubri", ""]
    )
    feuille_benef.append(
        ["Kaboré", "Issa", "M", "1985-11-03", "70000002", "CNIB654321", "CNIB", "BF", "Kadiogo", "Kadiogo", "Koubri", ""]
    )

    generer_feuille_paliers_alerte(wb)

    instructions = wb.create_sheet("Instructions")
    instructions.append(["Planification complète — consignes"])
    instructions.append([
        "Ce classeur crée en une fois : un cadre stratégique, un projet et sa planification, ses "
        "indicateurs et ses bénéficiaires (automatiquement inscrits à ce projet)."
    ])
    instructions.append(["- Seule la feuille « Projet » est obligatoire (Nom, Code, dates de début/fin)."])
    instructions.append([
        "- Feuille « Cadre stratégique » : laisse la ligne « Nom du cadre stratégique » vide si tu ne veux "
        "pas en créer un — le projet sera alors créé sans cadre stratégique."
    ])
    instructions.append([
        "- Feuille « Structuration » : une colonne = un niveau, une ligne = un élément dans UNE SEULE "
        "colonne (voir modèle du Plan stratégique pour le détail du format)."
    ])
    instructions.append([
        "- Feuille « Planification » : même principe pour Objectif général > Objectif spécifique > "
        "Activité > Sous-activité."
    ])
    instructions.append(["- Feuille « Indicateurs » : rattachés au projet créé dans ce même classeur."])
    instructions.append(["- Feuille « Bénéficiaires » : chaque bénéficiaire créé est automatiquement inscrit à ce projet."])
    instructions.append([
        "- Feuille « Grilles d'alerte » (facultative) : pour donner au projet, à un indicateur ou à une "
        "activité précis (par son libellé exact, créé dans ce même classeur) sa propre grille de paliers "
        "plutôt que d'hériter de la grille globale — une ligne par palier, plusieurs lignes pour la même "
        "portée forment sa grille complète."
    ])
    instructions.append([
        "- Pour importer séparément (un projet seul, des indicateurs seuls...), utilise les modèles dédiés "
        "depuis chaque écran (Projets, Indicateurs, Bénéficiaires, Cadre stratégique)."
    ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def importer_planification_complete(fichier, utilisateur) -> dict:
    from apps.beneficiaries.services.import_excel import importer_beneficiaires_depuis_feuille
    from apps.indicators.import_excel import importer_indicateurs_depuis_feuille
    from apps.projects.import_excel import _lire_champs_cle_valeur, _texte, importer_projet_depuis_wb
    from apps.projects.models import Projet
    from apps.strategy.models import CadreStrategique
    from apps.strategy.services import importer_structuration_depuis_feuille

    wb = load_workbook(fichier, data_only=True)

    resultat = {
        "cadre_strategique_id": None,
        "cadre_strategique_nom": None,
        "elements_strategiques_crees": 0,
        "projet_id": None,
        "projet_code": None,
        "objectifs_generaux": 0,
        "objectifs_specifiques_crees": 0,
        "activites_creees": 0,
        "sous_activites_creees": 0,
        "indicateurs_crees": 0,
        "beneficiaires_crees": 0,
        "doublons_detectes": 0,
        "grilles_alerte_creees": 0,
        "erreurs": [],
        "avertissements": [],
    }

    # ---------- 1. Cadre stratégique (facultatif) ----------
    cadre = None
    if "Cadre stratégique" in wb.sheetnames:
        champs_cadre = _lire_champs_cle_valeur(wb["Cadre stratégique"])
        nom_cadre = _texte(champs_cadre, "Nom du cadre stratégique")
        if nom_cadre:
            cadre = CadreStrategique.objects.create(nom=nom_cadre, description=_texte(champs_cadre, "Description"))
            resultat["cadre_strategique_id"] = cadre.id
            resultat["cadre_strategique_nom"] = cadre.nom
            if "Structuration" in wb.sheetnames:
                resultat_struct = importer_structuration_depuis_feuille(wb["Structuration"], cadre)
                resultat["elements_strategiques_crees"] = resultat_struct["elements_crees"]
                resultat["erreurs"] += _prefixer(resultat_struct["erreurs"], "Structuration")

    # ---------- 2. Projet + planification (obligatoire) ----------
    resultat_projet = importer_projet_depuis_wb(wb, cadre_strategique_impose=cadre)
    resultat["projet_id"] = resultat_projet.get("projet_id")
    resultat["projet_code"] = resultat_projet.get("projet_code")
    resultat["objectifs_generaux"] = resultat_projet.get("objectifs_generaux", 0)
    resultat["objectifs_specifiques_crees"] = resultat_projet.get("objectifs_specifiques_crees", 0)
    resultat["activites_creees"] = resultat_projet.get("activites_creees", 0)
    resultat["sous_activites_creees"] = resultat_projet.get("sous_activites_creees", 0)
    resultat["erreurs"] += _prefixer(resultat_projet["erreurs"], "Projet")
    resultat["avertissements"] += _prefixer(resultat_projet["avertissements"], "Projet")

    if not resultat_projet.get("projet_id"):
        # Pas de projet créé : indicateurs et bénéficiaires en dépendent, on s'arrête ici.
        return resultat

    projet = Projet.objects.get(id=resultat_projet["projet_id"])

    # ---------- 3. Indicateurs (facultatif) ----------
    indicateurs_par_nom = {}
    if "Indicateurs" in wb.sheetnames:
        resultat_indic = importer_indicateurs_depuis_feuille(wb["Indicateurs"], projet=projet)
        resultat["indicateurs_crees"] = resultat_indic["crees"]
        indicateurs_par_nom = resultat_indic.get("_indicateurs_par_nom", {})
        resultat["erreurs"] += _prefixer(resultat_indic["erreurs"], "Indicateurs")
        resultat["avertissements"] += _prefixer(resultat_indic["avertissements"], "Indicateurs")

    # ---------- 4. Bénéficiaires (facultatif) ----------
    if "Bénéficiaires" in wb.sheetnames:
        resultat_benef = importer_beneficiaires_depuis_feuille(wb["Bénéficiaires"], utilisateur, projet=projet)
        resultat["beneficiaires_crees"] = resultat_benef["crees"]
        resultat["doublons_detectes"] = resultat_benef["doublons_detectes"]
        resultat["erreurs"] += _prefixer(resultat_benef["erreurs"], "Bénéficiaires")
        resultat["avertissements"] += _prefixer(resultat_benef["avertissements"], "Bénéficiaires")

    # ---------- 5. Grilles d'alerte (facultatif) ----------
    if "Grilles d'alerte" in wb.sheetnames:
        from apps.indicators.import_excel import importer_paliers_alerte
        from apps.projects.models import Activite

        activites_par_nom = {
            a.libelle.strip().lower(): a
            for a in Activite.objects.filter(objectif_specifique__objectif_general__projet=projet)
        }
        resultat_paliers = importer_paliers_alerte(
            wb["Grilles d'alerte"], projet=projet, indicateurs_par_nom=indicateurs_par_nom, activites_par_nom=activites_par_nom
        )
        resultat["grilles_alerte_creees"] = resultat_paliers["grilles_creees"]
        resultat["erreurs"] += _prefixer(resultat_paliers["erreurs"], "Grilles d'alerte")

    return resultat
