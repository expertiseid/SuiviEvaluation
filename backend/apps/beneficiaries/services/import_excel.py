"""
Import en masse des bénéficiaires depuis un fichier Excel — évite la saisie
un par un pour de gros effectifs. Best-effort : une ligne en erreur est
ignorée sans bloquer l'import des autres, et tout est rapporté au client.
"""
import datetime
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from apps.core.excel_utils import figer_entete, quadriller

COLONNES = [
    "Nom",
    "Prénom",
    "Sexe (F/M)",
    "Date de naissance (AAAA-MM-JJ)",
    "Téléphone",
    "N° pièce d'identité",
    "Type de pièce",
    "Pays (code ISO, ex : BF)",
    "Région",
    "Province",
    "Commune / Village",
    "Statuts particuliers (séparés par ;)",
    "Types d'activité menée (séparés par ;)",
]


def generer_modele_import() -> bytes:
    wb = Workbook()
    feuille = wb.active
    feuille.title = "Bénéficiaires"
    feuille.append(COLONNES)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    feuille.append(
        [
            "Traoré",
            "Awa",
            "F",
            "1990-05-12",
            "70000001",
            "CNIB123456",
            "CNIB",
            "BF",
            "Kadiogo",
            "Kadiogo",
            "Koubri",
            "Femme;Jeune",
            "Maraîchage",
        ]
    )
    for _ in range(20):
        feuille.append([""] * len(COLONNES))
    quadriller(feuille, max_col=len(COLONNES))
    figer_entete(feuille)

    instructions = wb.create_sheet("Instructions")
    instructions.append(["Consignes"])
    instructions.append(["- Ne pas modifier l'ordre ou le nom des colonnes de la feuille 'Bénéficiaires'."])
    instructions.append(["- Nom, Prénom et Sexe sont obligatoires ; les autres colonnes sont facultatives."])
    instructions.append(["- Sexe : F ou M."])
    instructions.append(["- Pays : code ISO 3166-1 alpha-2 (ex : BF). Par défaut BF si laissé vide."])
    instructions.append([
        "- Région/Province/Commune doivent correspondre aux noms des niveaux administratifs configurés pour ce "
        "pays dans la plateforme (menu « Zones administratives »)."
    ])
    instructions.append(["- Commune / Village : texte libre, créé automatiquement s'il n'existe pas encore."])
    instructions.append(["- Plusieurs statuts particuliers : séparer par un point-virgule (ex : Femme;Jeune)."])
    instructions.append([
        "- Types d'activité menée (ex : Maraîchage, Élevage) : séparer par un point-virgule — un type inconnu "
        "est ignoré (avertissement), il doit d'abord être créé (fiche bénéficiaire ou admin)."
    ])
    instructions.append(["- Les doublons potentiels seront signalés automatiquement après l'import."])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _normaliser_sexe(valeur) -> str | None:
    if not valeur:
        return None
    v = str(valeur).strip().lower()
    if v in ("f", "féminin", "feminin", "femme"):
        return "F"
    if v in ("m", "masculin", "homme"):
        return "M"
    return None


def _normaliser_date(valeur):
    if not valeur:
        return None
    if isinstance(valeur, (datetime.date, datetime.datetime)):
        return valeur.date() if isinstance(valeur, datetime.datetime) else valeur
    try:
        return datetime.datetime.strptime(str(valeur).strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _resoudre_zone(pays, region_nom, province_nom, village_nom):
    """
    Résout Région/Province/Commune-Village en une Zone, en s'appuyant sur les
    niveaux administratifs configurés pour ce pays (nom_niveau « Région »,
    « Province »... insensible à la casse) — si un niveau de ce nom n'est pas
    configuré pour ce pays, la colonne correspondante est ignorée.
    """
    from apps.geo.models import NiveauAdministratif, Zone

    def niveau(nom_niveau):
        return NiveauAdministratif.objects.filter(pays=pays, nom_niveau__iexact=nom_niveau).first()

    zone_finale = None

    if region_nom and (niveau_region := niveau("Région")):
        zone_finale = Zone.objects.filter(
            niveau_administratif=niveau_region, nom__iexact=str(region_nom).strip()
        ).first()

    if province_nom and (niveau_province := niveau("Province")):
        qs = Zone.objects.filter(niveau_administratif=niveau_province, nom__iexact=str(province_nom).strip())
        if zone_finale:
            qs = qs.filter(parent=zone_finale)
        province = qs.first()
        if province:
            zone_finale = province

    if village_nom:
        niveau_village = niveau("Village") or niveau("Commune")
        if niveau_village:
            nom = str(village_nom).strip()
            village = Zone.objects.filter(
                niveau_administratif=niveau_village, nom__iexact=nom, parent=zone_finale
            ).first()
            if not village:
                village = Zone.objects.create(nom=nom, niveau_administratif=niveau_village, parent=zone_finale)
            zone_finale = village

    return zone_finale


def _resoudre_statuts(valeur):
    from ..models import StatutParticulier

    if not valeur:
        return [], []

    noms = [n.strip() for n in str(valeur).split(";") if n.strip()]
    trouves = []
    inconnus = []
    for nom in noms:
        statut = StatutParticulier.objects.filter(libelle__iexact=nom).first() or StatutParticulier.objects.filter(
            code__iexact=nom
        ).first()
        if statut:
            trouves.append(statut)
        else:
            inconnus.append(nom)
    return trouves, inconnus


def _resoudre_types_activite(valeur):
    from ..models import TypeActiviteBeneficiaire

    if not valeur:
        return [], []

    noms = [n.strip() for n in str(valeur).split(";") if n.strip()]
    trouves = []
    inconnus = []
    for nom in noms:
        type_activite = TypeActiviteBeneficiaire.objects.filter(
            libelle__iexact=nom
        ).first() or TypeActiviteBeneficiaire.objects.filter(code__iexact=nom).first()
        if type_activite:
            trouves.append(type_activite)
        else:
            inconnus.append(nom)
    return trouves, inconnus


def importer_beneficiaires(fichier, utilisateur, projet=None) -> dict:
    wb = load_workbook(fichier, data_only=True)
    feuille = wb["Bénéficiaires"] if "Bénéficiaires" in wb.sheetnames else wb.worksheets[0]
    return importer_beneficiaires_depuis_feuille(feuille, utilisateur, projet=projet)


def importer_beneficiaires_depuis_feuille(feuille, utilisateur, projet=None) -> dict:
    """
    `projet` : si fourni (import combiné), chaque bénéficiaire créé est
    aussi inscrit à ce projet (ParticipationProjet, date du jour).
    """
    from ..models import Beneficiaire, SignalementDoublon
    from .duplicate_detection import rechercher_doublons

    lignes = list(feuille.iter_rows(min_row=1, values_only=True))
    if not lignes:
        return {"crees": 0, "erreurs": [], "avertissements": [], "doublons_detectes": 0}

    entetes = [str(c).strip() if c else "" for c in lignes[0]]

    def valeur(ligne, nom_colonne):
        if nom_colonne not in entetes:
            return None
        return ligne[entetes.index(nom_colonne)]

    crees = 0
    erreurs = []
    avertissements = []
    doublons_detectes = 0

    for numero, ligne in enumerate(lignes[1:], start=2):
        if not any(ligne):
            continue

        nom = valeur(ligne, "Nom")
        prenom = valeur(ligne, "Prénom")
        sexe = _normaliser_sexe(valeur(ligne, "Sexe (F/M)"))

        if not nom or not prenom:
            erreurs.append({"ligne": numero, "message": "Nom et prénom obligatoires — ligne ignorée."})
            continue
        if not sexe:
            erreurs.append({"ligne": numero, "message": "Sexe invalide (attendu F ou M) — ligne ignorée."})
            continue

        pays = str(valeur(ligne, "Pays (code ISO, ex : BF)") or "BF").strip().upper() or "BF"
        zone = _resoudre_zone(
            pays, valeur(ligne, "Région"), valeur(ligne, "Province"), valeur(ligne, "Commune / Village")
        )
        if (valeur(ligne, "Région") or valeur(ligne, "Province")) and zone is None:
            avertissements.append({"ligne": numero, "message": "Région/Province non trouvée — zone laissée vide."})

        statuts, inconnus = _resoudre_statuts(valeur(ligne, "Statuts particuliers (séparés par ;)"))
        for nom_inconnu in inconnus:
            avertissements.append({"ligne": numero, "message": f"Statut particulier inconnu ignoré : « {nom_inconnu} »."})

        types_activite, types_inconnus = _resoudre_types_activite(valeur(ligne, "Types d'activité menée (séparés par ;)"))
        for nom_inconnu in types_inconnus:
            avertissements.append({"ligne": numero, "message": f"Type d'activité inconnu ignoré : « {nom_inconnu} »."})

        beneficiaire = Beneficiaire.objects.create(
            nom=str(nom).strip(),
            prenom=str(prenom).strip(),
            sexe=sexe,
            date_naissance=_normaliser_date(valeur(ligne, "Date de naissance (AAAA-MM-JJ)")),
            telephone=str(valeur(ligne, "Téléphone") or "").strip(),
            numero_piece_identite=str(valeur(ligne, "N° pièce d'identité") or "").strip(),
            type_piece=str(valeur(ligne, "Type de pièce") or "").strip(),
            pays=pays,
            zone=zone,
        )
        if statuts:
            beneficiaire.statuts_particuliers.set(statuts)
        if types_activite:
            beneficiaire.types_activite.set(types_activite)

        if projet is not None:
            from ..models import ParticipationProjet

            ParticipationProjet.objects.create(
                beneficiaire=beneficiaire, projet=projet, date_inscription=datetime.date.today()
            )

        for resultat in rechercher_doublons(beneficiaire):
            b1, b2 = sorted([beneficiaire, resultat["candidat"]], key=lambda b: b.pk)
            _, cree = SignalementDoublon.objects.get_or_create(
                beneficiaire_1=b1,
                beneficiaire_2=b2,
                defaults={"score": resultat["score"], "methode": resultat["methode"]},
            )
            if cree:
                doublons_detectes += 1

        crees += 1

    return {
        "crees": crees,
        "erreurs": erreurs,
        "avertissements": avertissements,
        "doublons_detectes": doublons_detectes,
    }
