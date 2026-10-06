"""
Import en masse de la structuration stratégique (niveaux + éléments) depuis
un fichier Excel — évite de tout ressaisir à la main écran par écran.

Format "plan" (une colonne par niveau, comme un sommaire Word) plutôt qu'un
système de Code/Parent à taper à la main : trop technique pour un
utilisateur non-développeur (confusion constatée en usage réel). La
hiérarchie se déduit uniquement de la position :
- une ligne ne remplit qu'une seule colonne (= un seul niveau) ;
- son parent est le dernier élément rencontré dans la colonne immédiatement
  à sa gauche (ou plus à gauche encore si le niveau autorise le saut de
  niveau).

Tout est scopé à un CadreStrategique précis : plusieurs cadres coexistent
sans jamais se mélanger (niveaux et éléments propres à chacun).

Best-effort : une ligne en erreur est ignorée sans bloquer l'import des
autres, et tout est rapporté au client. Rejouable : réimporter le même
fichier met à jour les éléments déjà créés (même niveau, même nom, même
parent) plutôt que de les dupliquer.
"""
import io
import re

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from apps.core.excel_utils import figer_entete, quadriller

# Reconnaît le format "CODE — Nom" tel que généré par exporter_structuration
# (code = un seul mot sans espace, ex : "AXE1", "OR1.1") — permet de
# réimporter tel quel un export sans dupliquer les éléments qui ont un code
# (le rattachement se ferait sinon sur "CODE — Nom" entier, jamais égal au
# nom seul stocké en base).
RE_CODE_NOM = re.compile(r"^(\S+) — (.+)$")


def _decouper_code_nom(texte: str) -> tuple[str, str]:
    match = RE_CODE_NOM.match(texte)
    if match:
        return match.group(1), match.group(2)
    return "", texte

NOMS_NIVEAUX_EXEMPLE = ["Orientation stratégique", "Axe d'intervention"]

LIGNES_EXEMPLE = [
    ["Amplifier la diffusion des pratiques agroécologiques et de protection de l'environnement", ""],
    ["", "Le renforcement des capacités des producteurs et productrices"],
    ["", "Le renforcement de la structuration des producteurs et productrices accompagnés"],
]


def generer_modele_import(cadre_strategique) -> bytes:
    from .models import TypeNiveau

    niveaux = list(TypeNiveau.objects.filter(cadre_strategique=cadre_strategique).order_by("ordre"))
    noms_colonnes = [n.nom_niveau for n in niveaux] if niveaux else NOMS_NIVEAUX_EXEMPLE

    wb = Workbook()
    feuille = wb.active
    feuille.title = "Structuration"
    feuille.append(noms_colonnes)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)

    if not niveaux:
        for ligne in LIGNES_EXEMPLE:
            feuille.append(ligne)

    for _ in range(20):
        feuille.append([""] * len(noms_colonnes))
    quadriller(feuille, max_col=len(noms_colonnes))
    figer_entete(feuille)
    for i in range(len(noms_colonnes)):
        feuille.column_dimensions[chr(ord("A") + i)].width = 42

    instructions = wb.create_sheet("Instructions")
    instructions.append([f"Cadre stratégique : {cadre_strategique.nom}"])
    instructions.append(["Consignes"])
    instructions.append(["- Une colonne = un niveau de la hiérarchie (dans l'ordre, de gauche à droite)."])
    instructions.append(["- Sur chaque ligne, écris le nom de l'élément dans UNE SEULE colonne : celle de son niveau."])
    instructions.append(["- Son \"parent\" est déduit automatiquement : c'est le dernier élément écrit dans la"])
    instructions.append(["  colonne juste à gauche, au-dessus de cette ligne."])
    instructions.append(["- Exemple : une ligne \"Axe d'intervention\" se rattache au dernier \"Orientation"])
    instructions.append(["  stratégique\" écrit au-dessus d'elle."])
    instructions.append(["- Si le nom d'une colonne n'existe pas encore comme niveau dans ce cadre, il est"])
    instructions.append(["  créé automatiquement lors de l'import."])
    instructions.append(["- Réimporter le même fichier met à jour les éléments déjà importés (même nom, même"])
    instructions.append(["  position) au lieu de les dupliquer."])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def exporter_structuration(cadre_strategique) -> bytes:
    """
    Export au même format "plan" que le modèle d'import (une colonne par
    niveau) mais rempli avec les éléments réels du cadre — le fichier généré
    est donc lui-même réimportable tel quel (voir importer_structuration).
    """
    from .models import ElementStrategique, TypeNiveau

    niveaux = list(TypeNiveau.objects.filter(cadre_strategique=cadre_strategique).order_by("ordre"))
    noms_colonnes = [n.nom_niveau for n in niveaux]
    colonne_par_niveau = {n.id: i for i, n in enumerate(niveaux)}

    elements = list(
        ElementStrategique.objects.filter(type_niveau__cadre_strategique=cadre_strategique).select_related(
            "type_niveau"
        )
    )
    enfants_par_parent: dict = {}
    for element in elements:
        enfants_par_parent.setdefault(element.element_parent_id, []).append(element)
    for liste in enfants_par_parent.values():
        liste.sort(key=lambda e: (e.code or "", e.nom))

    wb = Workbook()
    feuille = wb.active
    feuille.title = "Structuration"
    feuille.append(noms_colonnes)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for i in range(len(noms_colonnes)):
        feuille.column_dimensions[chr(ord("A") + i)].width = 42

    def visiter(element):
        colonne = colonne_par_niveau.get(element.type_niveau_id)
        if colonne is not None:
            ligne = [""] * len(noms_colonnes)
            ligne[colonne] = f"{element.code} — {element.nom}" if element.code else element.nom
            feuille.append(ligne)
        for enfant in enfants_par_parent.get(element.id, []):
            visiter(enfant)

    for racine in enfants_par_parent.get(None, []):
        visiter(racine)

    for _ in range(15):
        feuille.append([""] * len(noms_colonnes))
    quadriller(feuille, max_col=len(noms_colonnes))
    figer_entete(feuille)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _resoudre_type_niveau(nom_niveau, cache, cadre_strategique):
    """Réutilise un TypeNiveau existant de ce cadre (par nom, insensible à la
    casse) ou en crée un nouveau en fin de chaîne — même comportement que le
    bouton "Ajouter un niveau" de l'écran de configuration."""
    from .models import TypeNiveau

    cle = nom_niveau.strip().lower()
    if cle in cache:
        return cache[cle]

    existant = TypeNiveau.objects.filter(
        cadre_strategique=cadre_strategique, nom_niveau__iexact=nom_niveau.strip()
    ).first()
    if existant:
        cache[cle] = existant
        return existant

    dernier = TypeNiveau.objects.filter(cadre_strategique=cadre_strategique).order_by("-ordre").first()
    niveau = TypeNiveau.objects.create(
        cadre_strategique=cadre_strategique,
        nom_niveau=nom_niveau.strip(),
        ordre=(dernier.ordre if dernier else 0) + 1,
        niveau_parent=dernier,
    )
    cache[cle] = niveau
    return niveau


def importer_structuration(fichier, cadre_strategique) -> dict:
    wb = load_workbook(fichier, data_only=True)
    feuille = wb["Structuration"] if "Structuration" in wb.sheetnames else wb.worksheets[0]
    return importer_structuration_depuis_feuille(feuille, cadre_strategique)


def importer_structuration_depuis_feuille(feuille, cadre_strategique) -> dict:
    from .models import ElementStrategique, TypeNiveau

    lignes = list(feuille.iter_rows(min_row=1, values_only=True))
    if not lignes:
        return {"niveaux_crees": 0, "elements_crees": 0, "elements_mis_a_jour": 0, "erreurs": []}

    noms_colonnes = [str(c).strip() if c else "" for c in lignes[0]]
    while noms_colonnes and not noms_colonnes[-1]:
        noms_colonnes.pop()

    if not noms_colonnes:
        return {
            "niveaux_crees": 0,
            "elements_crees": 0,
            "elements_mis_a_jour": 0,
            "erreurs": [{"ligne": 1, "message": "Aucune colonne de niveau trouvée sur la première ligne."}],
        }

    niveaux_avant = TypeNiveau.objects.filter(cadre_strategique=cadre_strategique).count()
    cache_niveaux: dict[str, TypeNiveau] = {}
    types_par_colonne = [_resoudre_type_niveau(nom, cache_niveaux, cadre_strategique) for nom in noms_colonnes]

    elements_existants: dict[tuple[int, str, int | None], ElementStrategique] = {
        (e.type_niveau_id, e.nom.strip().lower(), e.element_parent_id): e
        for e in ElementStrategique.objects.filter(type_niveau__cadre_strategique=cadre_strategique)
    }
    dernier_par_colonne: list[ElementStrategique | None] = [None] * len(noms_colonnes)
    elements_crees = 0
    elements_mis_a_jour = 0
    erreurs = []

    for numero, ligne in enumerate(lignes[1:], start=2):
        cellules = list(ligne[: len(noms_colonnes)])
        colonnes_remplies = [i for i, v in enumerate(cellules) if v not in (None, "")]

        if not colonnes_remplies:
            continue
        if len(colonnes_remplies) > 1:
            erreurs.append(
                {"ligne": numero, "message": "Une seule colonne doit être remplie par ligne — ligne ignorée."}
            )
            continue

        idx = colonnes_remplies[0]
        code, nom = _decouper_code_nom(str(cellules[idx]).strip())
        type_niveau = types_par_colonne[idx]

        element_parent = None
        if idx > 0:
            if dernier_par_colonne[idx - 1] is not None:
                element_parent = dernier_par_colonne[idx - 1]
            elif type_niveau.saut_niveau_autorise:
                for j in range(idx - 2, -1, -1):
                    if dernier_par_colonne[j] is not None:
                        element_parent = dernier_par_colonne[j]
                        break
            if element_parent is None:
                erreurs.append(
                    {
                        "ligne": numero,
                        "message": f"Aucun élément « {noms_colonnes[idx - 1]} » renseigné au-dessus pour rattacher cette ligne — ligne ignorée.",
                    }
                )
                continue

        cle = (type_niveau.id, nom.lower(), element_parent.id if element_parent else None)
        existant = elements_existants.get(cle)
        instance = existant or ElementStrategique(type_niveau=type_niveau, code="")
        instance.type_niveau = type_niveau
        instance.element_parent = element_parent
        instance.nom = nom
        if code:
            instance.code = code

        try:
            instance.clean()
        except DjangoValidationError as exc:
            erreurs.append({"ligne": numero, "message": " ".join(exc.messages)})
            continue

        instance.save()
        elements_existants[cle] = instance
        dernier_par_colonne[idx] = instance
        for j in range(idx + 1, len(dernier_par_colonne)):
            dernier_par_colonne[j] = None

        if existant:
            elements_mis_a_jour += 1
        else:
            elements_crees += 1

    return {
        "niveaux_crees": TypeNiveau.objects.filter(cadre_strategique=cadre_strategique).count() - niveaux_avant,
        "elements_crees": elements_crees,
        "elements_mis_a_jour": elements_mis_a_jour,
        "erreurs": erreurs,
    }


def supprimer_cadre_strategique_cascade(cadre):
    """
    TypeNiveau/ElementStrategique sont en PROTECT (même logique que Projet :
    ne jamais perdre une arborescence par accident via un autre chemin de
    code), donc la suppression volontaire d'un cadre stratégique doit
    désimbriquer manuellement, des feuilles vers la racine, dans une
    transaction — même pattern que apps.projects.services.supprimer_projet_cascade.

    Les éléments sont supprimés feuille par feuille en suivant leur
    element_parent réel plutôt que le TypeNiveau.ordre attendu : un
    réordonnancement des niveaux après coup peut désynchroniser les deux,
    et supprimer par ordre supposé lèverait alors un ProtectedError.
    """
    from .models import ElementStrategique, TypeNiveau

    with transaction.atomic():
        restants = {e.id: e for e in ElementStrategique.objects.filter(type_niveau__cadre_strategique=cadre)}
        while restants:
            parents_references = {e.element_parent_id for e in restants.values() if e.element_parent_id}
            feuilles = [eid for eid in restants if eid not in parents_references]
            if not feuilles:
                break
            ElementStrategique.objects.filter(id__in=feuilles).delete()
            for eid in feuilles:
                restants.pop(eid)

        niveaux_desc = list(TypeNiveau.objects.filter(cadre_strategique=cadre).order_by("-ordre"))
        for niveau in niveaux_desc:
            niveau.delete()
        cadre.delete()
