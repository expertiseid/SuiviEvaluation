"""
Export « bien fait » de la fiche de suivi d'un projet : reprend TOUTES les
informations affichées sur le tableau de bord de suivi (exécution globale,
indicateurs, activités et sous-activités avec leur historique complet et la
progression attendue à date), réparties sur plusieurs feuilles lisibles et
filtrables — plutôt qu'un import, ce fichier est un document de restitution
(bailleur, rapport, archive).
"""
import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .services import construire_dashboard_suivi, ecart_texte, taux_de, valeur_attendue_a

COULEUR_ENTETE = "1F4E78"
COULEUR_TITRE = "D9E2F3"

STATUT_ACTIVITE_LABEL = {"NON_REALISEE": "Non réalisée", "EN_COURS": "En cours", "REALISEE": "Réalisée"}
STATUT_SOUS_ACTIVITE_LABEL = {"PLANIFIEE": "Planifiée", "EN_COURS": "En cours", "TERMINEE": "Terminée"}


def _entete(feuille, colonnes, ligne=1):
    feuille.append(colonnes)
    for cellule in feuille[ligne]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor=COULEUR_ENTETE)
        cellule.alignment = Alignment(vertical="center", wrap_text=True)
    # Une référence texte plutôt que feuille.cell(...) : cette dernière CRÉE
    # la cellule visée, ce qui décale ws.max_row et fait sauter une ligne au
    # prochain .append() — piège classique d'openpyxl.
    feuille.freeze_panes = f"A{ligne + 1}"
    feuille.auto_filter.ref = f"A{ligne}:{get_column_letter(len(colonnes))}{ligne}"


def _largeurs(feuille, largeurs):
    for i, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(i)].width = largeur


def exporter_suivi_excel(projet) -> bytes:
    data = construire_dashboard_suivi(projet)
    aujourdhui = datetime.date.today()

    wb = Workbook()

    # ---------- Résumé ----------
    resume = wb.active
    resume.title = "Résumé"
    resume.append([f"Suivi — {data['projet']['code']} — {data['projet']['nom']}"])
    resume.cell(row=1, column=1).font = Font(bold=True, size=14)
    resume.append([f"Généré le {aujourdhui.strftime('%d/%m/%Y')}"])
    resume.cell(row=2, column=1).font = Font(italic=True, color="666666")
    resume.append([])
    _entete(resume, ["Indicateur global", "Valeur"], ligne=4)
    physique = data["taux_execution_physique_global"]
    financier = data["taux_execution_financiere_global"]
    statut_global_label = {"EN_RETARD": "En retard", "ATTEINT": "Atteint", "EN_COURS": "En cours"}.get(
        data["statut_global"], data["statut_global"]
    )
    resume.append(["Exécution physique globale (%)", physique if physique is not None else "—"])
    resume.append(["Exécution financière globale (%)", financier if financier is not None else "—"])
    resume.append(["Statut global", statut_global_label])
    resume.append(["Nombre d'indicateurs", len(data["indicateurs"])])
    resume.append(["Nombre d'activités", len(data["activites"])])
    _largeurs(resume, [34, 20])

    # ---------- Indicateurs (résumé) ----------
    ws_indic = wb.create_sheet("Indicateurs")
    colonnes = [
        "Libellé",
        "Unité",
        "Valeur de référence",
        "Cible",
        "Réalisé cumulé",
        "Taux (%)",
        "Statut",
        "Planifié début",
        "Planifié fin",
        "Attendu aujourd'hui",
        "Taux attendu (%)",
        "Écart",
    ]
    _entete(ws_indic, colonnes)
    for indicateur in data["indicateurs"]:
        cible = indicateur["valeur_cible"]
        depart = indicateur["valeur_reference"] or 0
        debut_p = indicateur["periode_debut_planifiee"]
        fin_p = indicateur["periode_fin_planifiee"]
        realise = indicateur["historique"][-1]["valeur_cumulee"] if indicateur["historique"] else None
        attendu = valeur_attendue_a(aujourdhui, debut_p, fin_p, depart, cible) if debut_p and fin_p else None
        ws_indic.append(
            [
                indicateur["libelle"],
                indicateur["unite"],
                float(depart),
                float(cible),
                float(realise) if realise is not None else None,
                indicateur["taux_actuel"],
                indicateur["palier_actuel"]["libelle"] if indicateur["palier_actuel"] else "—",
                debut_p,
                fin_p,
                round(attendu, 1) if attendu is not None else None,
                taux_de(attendu, cible) if attendu is not None else None,
                ecart_texte(realise, attendu, indicateur["unite"]),
            ]
        )
    _largeurs(ws_indic, [38, 10, 14, 12, 14, 10, 14, 14, 14, 16, 14, 38])

    # ---------- Indicateurs — historique ----------
    ws_indic_hist = wb.create_sheet("Indicateurs — historique")
    colonnes = [
        "Indicateur",
        "Période début",
        "Période fin",
        "Valeur réalisée",
        "Cumul",
        "Taux cumul (%)",
        "Attendu",
        "Taux attendu (%)",
        "Écart",
        "Commentaire",
    ]
    _entete(ws_indic_hist, colonnes)
    for indicateur in data["indicateurs"]:
        cible = indicateur["valeur_cible"]
        depart = indicateur["valeur_reference"] or 0
        debut_p = indicateur["periode_debut_planifiee"]
        fin_p = indicateur["periode_fin_planifiee"]
        for h in indicateur["historique"]:
            attendu = valeur_attendue_a(h["periode_fin"], debut_p, fin_p, depart, cible) if debut_p and fin_p else None
            ws_indic_hist.append(
                [
                    indicateur["libelle"],
                    h["periode_debut"],
                    h["periode_fin"],
                    float(h["valeur_realisee"]),
                    float(h["valeur_cumulee"]),
                    h["taux"],
                    round(attendu, 1) if attendu is not None else None,
                    taux_de(attendu, cible) if attendu is not None else None,
                    ecart_texte(h["valeur_cumulee"], attendu, indicateur["unite"]),
                    h["commentaire"],
                ]
            )
    _largeurs(ws_indic_hist, [38, 13, 13, 13, 12, 12, 12, 14, 38, 30])

    # ---------- Activités (résumé) ----------
    ws_act = wb.create_sheet("Activités")
    colonnes = [
        "Code",
        "Libellé",
        "Statut",
        "Début",
        "Fin",
        "Valeur de référence",
        "Quantité prévue",
        "Unité",
        "Réalisé cumulé",
        "Taux physique (%)",
        "Attendu physique aujourd'hui",
        "Budget alloué (FCFA)",
        "Budget réalisé cumulé (FCFA)",
        "Taux financier (%)",
        "Attendu financier aujourd'hui (FCFA)",
    ]
    _entete(ws_act, colonnes)
    for activite in data["activites"]:
        cible = activite["quantite_prevue"]
        depart = activite["valeur_reference"] or 0
        debut = activite["date_debut"]
        fin = activite["date_fin"]
        budget = activite["budget_alloue"]
        realise = activite["historique"][-1]["quantite_cumulee"] if activite["historique"] else None
        budget_realise = activite["historique"][-1]["budget_cumule"] if activite["historique"] else None
        attendu_phys = valeur_attendue_a(aujourdhui, debut, fin, depart, cible) if debut and fin and cible else None
        attendu_fin = valeur_attendue_a(aujourdhui, debut, fin, 0, budget) if debut and fin and budget else None
        ws_act.append(
            [
                activite["code_activite"],
                activite["libelle"],
                STATUT_ACTIVITE_LABEL.get(activite["statut"], activite["statut"]),
                debut,
                fin,
                float(depart),
                float(cible) if cible is not None else None,
                activite["unite_quantite"],
                float(realise) if realise is not None else None,
                activite["taux_realisation"],
                round(attendu_phys, 1) if attendu_phys is not None else None,
                float(budget) if budget is not None else None,
                float(budget_realise) if budget_realise is not None else None,
                activite["taux_execution_financiere"],
                round(attendu_fin, 1) if attendu_fin is not None else None,
            ]
        )
    _largeurs(ws_act, [12, 38, 14, 12, 12, 14, 14, 10, 14, 14, 18, 16, 18, 14, 18])

    # ---------- Activités — historique ----------
    ws_act_hist = wb.create_sheet("Activités — historique")
    colonnes = [
        "Activité",
        "Période début",
        "Période fin",
        "Quantité réalisée",
        "Cumul quantité",
        "Taux physique (%)",
        "Attendu physique",
        "Écart physique",
        "Budget réalisé (FCFA)",
        "Cumul budget (FCFA)",
        "Taux financier (%)",
        "Attendu financier (FCFA)",
        "Écart financier",
        "Statut",
        "Commentaire",
    ]
    _entete(ws_act_hist, colonnes)
    for activite in data["activites"]:
        cible = activite["quantite_prevue"]
        depart = activite["valeur_reference"] or 0
        debut = activite["date_debut"]
        fin = activite["date_fin"]
        budget = activite["budget_alloue"]
        for h in activite["historique"]:
            cumule = h["quantite_cumulee"]
            attendu = valeur_attendue_a(h["periode_fin"], debut, fin, depart, cible) if debut and fin and cible else None
            budget_cumule = h["budget_cumule"]
            attendu_budget = (
                valeur_attendue_a(h["periode_fin"], debut, fin, 0, budget) if debut and fin and budget else None
            )
            ws_act_hist.append(
                [
                    activite["libelle"],
                    h["periode_debut"],
                    h["periode_fin"],
                    float(h["quantite_realisee"]) if h["quantite_realisee"] is not None else None,
                    float(cumule) if cumule is not None else None,
                    taux_de(cumule, cible),
                    round(attendu, 1) if attendu is not None else None,
                    ecart_texte(cumule, attendu, activite["unite_quantite"]),
                    float(h["budget_realise"]) if h["budget_realise"] is not None else None,
                    float(budget_cumule) if budget_cumule is not None else None,
                    taux_de(budget_cumule, budget),
                    round(attendu_budget, 1) if attendu_budget is not None else None,
                    ecart_texte(budget_cumule, attendu_budget, "FCFA"),
                    STATUT_ACTIVITE_LABEL.get(h["statut"], h["statut"]),
                    h["commentaire"],
                ]
            )
    _largeurs(ws_act_hist, [30, 13, 13, 14, 12, 14, 14, 34, 16, 14, 14, 16, 34, 14, 30])

    # ---------- Sous-activités (résumé + historique) ----------
    ws_sa = wb.create_sheet("Sous-activités")
    colonnes = ["Activité parente", "Sous-activité", "Statut", "Début", "Fin", "Quantité prévue", "Unité", "Réalisé cumulé", "Taux (%)", "Attendu aujourd'hui"]
    _entete(ws_sa, colonnes)
    for activite in data["activites"]:
        for sa in activite["sous_activites"]:
            cible = sa["quantite_prevue"]
            debut = sa["date_debut"]
            fin = sa["date_fin"]
            realise = sa["historique"][-1]["quantite_cumulee"] if sa["historique"] else None
            taux = taux_de(realise, cible)
            attendu = valeur_attendue_a(aujourdhui, debut, fin, 0, cible) if debut and fin and cible else None
            ws_sa.append(
                [
                    activite["libelle"],
                    sa["libelle"],
                    STATUT_SOUS_ACTIVITE_LABEL.get(sa["statut"], sa["statut"]),
                    debut,
                    fin,
                    float(cible) if cible is not None else None,
                    sa["unite_quantite"],
                    float(realise) if realise is not None else None,
                    taux,
                    round(attendu, 1) if attendu is not None else None,
                ]
            )
    _largeurs(ws_sa, [30, 30, 14, 12, 12, 14, 10, 14, 10, 16])

    ws_sa_hist = wb.create_sheet("Sous-activités — historique")
    colonnes = ["Activité parente", "Sous-activité", "Période début", "Période fin", "Quantité réalisée", "Cumul", "Taux (%)", "Attendu", "Écart", "Statut", "Commentaire"]
    _entete(ws_sa_hist, colonnes)
    for activite in data["activites"]:
        for sa in activite["sous_activites"]:
            cible = sa["quantite_prevue"]
            debut = sa["date_debut"]
            fin = sa["date_fin"]
            for h in sa["historique"]:
                cumule = h["quantite_cumulee"]
                attendu = valeur_attendue_a(h["periode_fin"], debut, fin, 0, cible) if debut and fin and cible else None
                ws_sa_hist.append(
                    [
                        activite["libelle"],
                        sa["libelle"],
                        h["periode_debut"],
                        h["periode_fin"],
                        float(h["quantite_realisee"]) if h["quantite_realisee"] is not None else None,
                        float(cumule) if cumule is not None else None,
                        taux_de(cumule, cible),
                        round(attendu, 1) if attendu is not None else None,
                        ecart_texte(cumule, attendu, sa["unite_quantite"]),
                        STATUT_SOUS_ACTIVITE_LABEL.get(h["statut"], h["statut"]),
                        h["commentaire"],
                    ]
                )
    _largeurs(ws_sa_hist, [26, 26, 13, 13, 14, 12, 10, 12, 34, 14, 30])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
