"""
Import en masse d'indicateurs depuis un fichier Excel — même philosophie
« best-effort » que les autres imports (structuration, projet, bénéficiaires) :
une ligne en erreur est ignorée sans bloquer l'import des autres.
"""
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from apps.core.excel_utils import figer_entete, quadriller

COLONNES = [
    "Libellé",
    "Unité",
    "Valeur de référence",
    "Valeur cible",
    "Fréquence de collecte (MENSUELLE/TRIMESTRIELLE/SEMESTRIELLE/ANNUELLE)",
    "Code projet",
    "Niveau de rattachement (Aucun/Projet/Objectif général/Objectif spécifique/Activité)",
    "Nom de l'élément rattaché",
]

LIGNE_EXEMPLE = [
    # Distinct de l'activité liée : l'activité mesure l'extrant (combien de
    # personnes formées), l'indicateur mesure le résultat (l'effet de cette
    # formation) — même si les deux sont suivis dans le temps, leur « sens »
    # (direction) diffère volontairement pour ne pas se dupliquer.
    "Pourcentage de ménages ayant adopté les techniques agroécologiques",
    "%",
    0,
    60,
    "TRIMESTRIELLE",
    "PRJ-2026-001",
    "Activité",
    "Former les producteurs aux techniques agroécologiques",
]


def _remplir_feuille_indicateurs(feuille, avec_colonne_projet: bool = True):
    colonnes = COLONNES if avec_colonne_projet else [c for c in COLONNES if c != "Code projet"]
    feuille.append(colonnes)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    ligne_exemple = LIGNE_EXEMPLE if avec_colonne_projet else [v for c, v in zip(COLONNES, LIGNE_EXEMPLE) if c != "Code projet"]
    feuille.append(ligne_exemple)
    for _ in range(20):
        feuille.append([""] * len(colonnes))
    quadriller(feuille, max_col=len(colonnes))
    figer_entete(feuille)
    return feuille


def generer_feuille_indicateurs(wb, avec_colonne_projet: bool = True):
    """Ajoute une feuille « Indicateurs » à un classeur existant (utilisé par l'import combiné)."""
    return _remplir_feuille_indicateurs(wb.create_sheet("Indicateurs"), avec_colonne_projet=avec_colonne_projet)


def generer_modele_import_indicateurs(avec_colonne_projet: bool = True) -> bytes:
    wb = Workbook()
    feuille = wb.active
    feuille.title = "Indicateurs"
    _remplir_feuille_indicateurs(feuille, avec_colonne_projet=avec_colonne_projet)

    feuille_paliers = generer_feuille_paliers_alerte(wb, avec_exemple=False, avec_lignes_vides=False)
    feuille_paliers.append(["Indicateur", LIGNE_EXEMPLE[0], 0, "Faible", "#d03b3b"])
    feuille_paliers.append(["Indicateur", LIGNE_EXEMPLE[0], 60, "Bon", "#0ca30c"])
    for _ in range(15):
        feuille_paliers.append([""] * len(COLONNES_PALIERS))
    quadriller(feuille_paliers, max_col=len(COLONNES_PALIERS))

    instructions = wb.create_sheet("Instructions")
    instructions.append(["Consignes"])
    instructions.append(["- Libellé, Unité, Valeur cible et Fréquence de collecte sont obligatoires."])
    if avec_colonne_projet:
        instructions.append([
            "- Code projet : requis dès qu'un niveau de rattachement autre que « Aucun » est utilisé — doit "
            "correspondre au code exact d'un projet déjà créé."
        ])
    instructions.append([
        "- Niveau de rattachement « Projet » : laisse « Nom de l'élément rattaché » vide. Pour Objectif "
        "général/spécifique/Activité, indique le libellé exact de l'élément dans ce projet."
    ])
    instructions.append(["- Niveau « Aucun » (ou vide) : indicateur stratégique, sans rattachement projet."])
    instructions.append([
        "- L'indicateur suit la grille d'alerte (paliers/libellés/couleurs) du projet auquel il est "
        "rattaché — configurable depuis Paramètres d'alerte, pour chaque projet."
    ])
    instructions.append([
        "- Feuille « Grilles d'alerte » (facultative) : pour donner à UN indicateur précis (par son libellé "
        "exact, créé dans cette même feuille « Indicateurs ») sa propre grille de paliers, plutôt que "
        "d'hériter de celle de son projet — une ligne par palier, plusieurs lignes pour le même indicateur "
        "forment sa grille complète."
    ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


COLONNES_PALIERS = [
    "Portée (Projet/Indicateur/Activité)",
    "Nom de l'élément (laisser vide pour Projet)",
    "Seuil (%)",
    "Libellé",
    "Couleur (hexadécimale, ex : #d03b3b)",
]

LIGNES_EXEMPLE_PALIERS = [
    ["Projet", "", 0, "Critique", "#d03b3b"],
    ["Projet", "", 50, "À risque", "#fab219"],
    ["Projet", "", 80, "Conforme", "#0ca30c"],
]


def generer_feuille_paliers_alerte(wb, avec_exemple: bool = True, avec_lignes_vides: bool = True):
    """
    Ajoute une feuille « Grilles d'alerte » à un classeur existant — une
    ligne par palier, plusieurs lignes pour une même portée (Portée + Nom
    de l'élément identiques) forment sa grille complète. Portée « Projet » :
    laisse « Nom de l'élément » vide (c'est le projet créé dans ce même
    classeur). Portée « Indicateur »/« Activité » : indique le libellé exact
    d'un indicateur/une activité de ce même classeur (ou déjà existant dans
    le projet, pour un import qui ne fait qu'ajouter des grilles).

    `avec_lignes_vides=False` : à utiliser quand l'appelant ajoute encore des
    lignes après cet appel (le quadrillage doit alors être fait par
    l'appelant une fois toutes les lignes ajoutées, sinon les lignes vides
    pré-remplies se retrouveraient AVANT celles ajoutées ensuite).
    """
    feuille = wb.create_sheet("Grilles d'alerte")
    feuille.append(COLONNES_PALIERS)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    figer_entete(feuille)
    if avec_exemple:
        for ligne in LIGNES_EXEMPLE_PALIERS:
            feuille.append(ligne)
    feuille.column_dimensions["A"].width = 22
    feuille.column_dimensions["B"].width = 45
    feuille.column_dimensions["C"].width = 10
    feuille.column_dimensions["D"].width = 20
    feuille.column_dimensions["E"].width = 28
    if avec_lignes_vides:
        for _ in range(15):
            feuille.append([""] * len(COLONNES_PALIERS))
        quadriller(feuille, max_col=len(COLONNES_PALIERS))
    return feuille


def importer_paliers_alerte(feuille, *, projet=None, indicateurs_par_nom=None, activites_par_nom=None) -> dict:
    """
    Importe la feuille « Grilles d'alerte » — best-effort, une grille
    (groupe de lignes de même portée+nom) en erreur est ignorée sans
    bloquer les autres. Remplace intégralement la grille propre existante
    de chaque portée rencontrée (mêmes règles que l'écran Paramètres
    d'alerte : une grille vide reviendrait à hériter du niveau supérieur,
    mais ici une portée sans ligne n'est simplement pas touchée).
    """
    from .models import PalierAlerte

    resultat = {"grilles_creees": 0, "erreurs": []}
    lignes = list(feuille.iter_rows(min_row=1, values_only=True))
    if not lignes:
        return resultat

    entetes = [str(c).strip() if c else "" for c in lignes[0]]

    def valeur(ligne, nom_colonne):
        if nom_colonne not in entetes:
            return None
        idx = entetes.index(nom_colonne)
        return ligne[idx] if idx < len(ligne) else None

    groupes: dict[tuple[str, str], list[tuple[int, str, str]]] = {}
    for numero, ligne in enumerate(lignes[1:], start=2):
        if not any(ligne):
            continue
        portee = str(valeur(ligne, "Portée (Projet/Indicateur/Activité)") or "").strip().upper()
        nom = str(valeur(ligne, "Nom de l'élément (laisser vide pour Projet)") or "").strip()
        borne = valeur(ligne, "Seuil (%)")
        libelle = valeur(ligne, "Libellé")
        couleur = str(valeur(ligne, "Couleur (hexadécimale, ex : #d03b3b)") or "#495057").strip()

        if portee not in ("PROJET", "INDICATEUR", "ACTIVITÉ", "ACTIVITE"):
            resultat["erreurs"].append(
                {"ligne": numero, "message": f"Portée invalide « {portee} » (Projet/Indicateur/Activité) — ligne ignorée."}
            )
            continue
        if borne in (None, "") or not str(libelle or "").strip():
            resultat["erreurs"].append({"ligne": numero, "message": "Seuil (%) et Libellé sont obligatoires — ligne ignorée."})
            continue
        try:
            borne_int = int(borne)
        except (TypeError, ValueError):
            resultat["erreurs"].append({"ligne": numero, "message": f"Seuil « {borne} » invalide — ligne ignorée."})
            continue
        if not (0 <= borne_int <= 100):
            resultat["erreurs"].append({"ligne": numero, "message": f"Seuil « {borne_int} » hors de 0-100 — ligne ignorée."})
            continue

        groupes.setdefault((portee, nom.lower()), []).append((borne_int, str(libelle).strip(), couleur))

    for (portee, nom_lower), paliers in groupes.items():
        if portee == "PROJET":
            if projet is None:
                resultat["erreurs"].append({"ligne": 0, "message": "Grille « Projet » ignorée : aucun projet dans cet import."})
                continue
            cible = {"projet": projet}
        elif portee == "INDICATEUR":
            indicateur = (indicateurs_par_nom or {}).get(nom_lower)
            if indicateur is None:
                resultat["erreurs"].append(
                    {"ligne": 0, "message": f"Indicateur « {nom_lower} » introuvable — grille ignorée."}
                )
                continue
            cible = {"indicateur": indicateur}
        else:
            activite = (activites_par_nom or {}).get(nom_lower)
            if activite is None:
                resultat["erreurs"].append(
                    {"ligne": 0, "message": f"Activité « {nom_lower} » introuvable — grille ignorée."}
                )
                continue
            cible = {"activite": activite}

        if len(set(b for b, _, _ in paliers)) != len(paliers):
            resultat["erreurs"].append({"ligne": 0, "message": f"Seuils dupliqués pour « {nom_lower or 'Projet'} » — grille ignorée."})
            continue

        PalierAlerte.objects.filter(**cible).delete()
        PalierAlerte.objects.bulk_create(
            [PalierAlerte(**cible, borne_min=b, libelle=l, couleur=c) for b, l, c in paliers]
        )
        resultat["grilles_creees"] += 1

    return resultat


def _resoudre_rattachement(projet, niveau: str, nom_element: str):
    from apps.projects.models import Activite, ObjectifGeneral, ObjectifSpecifique

    niveau = (niveau or "").strip().upper()
    if niveau in ("", "AUCUN"):
        return {}
    if niveau == "PROJET":
        return {"projet": projet}

    nom_element = (nom_element or "").strip()
    if niveau in ("OBJECTIF GÉNÉRAL", "OBJECTIF GENERAL", "OG"):
        og = ObjectifGeneral.objects.filter(projet=projet).first()
        return {"objectif_general": og} if og else None
    if niveau in ("OBJECTIF SPÉCIFIQUE", "OBJECTIF SPECIFIQUE", "OS"):
        objectif_specifique = ObjectifSpecifique.objects.filter(
            objectif_general__projet=projet, libelle__iexact=nom_element
        ).first()
        return {"objectif_specifique": objectif_specifique} if objectif_specifique else None
    if niveau in ("ACTIVITÉ", "ACTIVITE"):
        activite = Activite.objects.filter(
            objectif_specifique__objectif_general__projet=projet, libelle__iexact=nom_element
        ).first()
        return {"activite": activite} if activite else None
    return None


def importer_indicateurs(fichier, projet=None) -> dict:
    """
    `projet` : si fourni (import combiné), tous les rattachements se font
    dans ce projet et la colonne « Code projet » est ignorée. Sinon
    (import autonome), chaque ligne doit préciser son propre code projet
    dès qu'elle utilise un rattachement.
    """
    wb = load_workbook(fichier, data_only=True)
    feuille = wb["Indicateurs"] if "Indicateurs" in wb.sheetnames else wb.worksheets[0]
    resultat = importer_indicateurs_depuis_feuille(feuille, projet=projet)
    if "Grilles d'alerte" in wb.sheetnames:
        resultat_paliers = importer_paliers_alerte(
            wb["Grilles d'alerte"], projet=projet, indicateurs_par_nom=resultat.pop("_indicateurs_par_nom", {})
        )
        resultat["grilles_alerte_creees"] = resultat_paliers["grilles_creees"]
        resultat["erreurs"] += [{"ligne": e["ligne"], "message": f"[Grilles d'alerte] {e['message']}"} for e in resultat_paliers["erreurs"]]
    else:
        resultat.pop("_indicateurs_par_nom", None)
    return resultat


def importer_indicateurs_depuis_feuille(feuille, projet=None) -> dict:
    from apps.projects.models import Projet

    from .models import Indicateur

    lignes = list(feuille.iter_rows(min_row=1, values_only=True))
    if not lignes:
        return {"crees": 0, "erreurs": [], "avertissements": [], "_indicateurs_par_nom": {}}

    entetes = [str(c).strip() if c else "" for c in lignes[0]]

    def valeur(ligne, nom_colonne):
        if nom_colonne not in entetes:
            return None
        idx = entetes.index(nom_colonne)
        return ligne[idx] if idx < len(ligne) else None

    frequences_valides = dict(Indicateur.Frequence.choices)
    crees = 0
    erreurs = []
    avertissements = []
    indicateurs_par_nom = {}

    for numero, ligne in enumerate(lignes[1:], start=2):
        if not any(ligne):
            continue

        libelle = valeur(ligne, "Libellé")
        unite = valeur(ligne, "Unité")
        cible = valeur(ligne, "Valeur cible")
        frequence = str(
            valeur(ligne, "Fréquence de collecte (MENSUELLE/TRIMESTRIELLE/SEMESTRIELLE/ANNUELLE)") or ""
        ).strip().upper()

        if not libelle or not unite or cible in (None, ""):
            erreurs.append({"ligne": numero, "message": "Libellé, unité et valeur cible sont obligatoires — ligne ignorée."})
            continue
        if frequence not in frequences_valides:
            erreurs.append({"ligne": numero, "message": f"Fréquence de collecte invalide « {frequence} » — ligne ignorée."})
            continue

        projet_ligne = projet
        if projet_ligne is None:
            code_projet = str(valeur(ligne, "Code projet") or "").strip()
            if code_projet:
                projet_ligne = Projet.objects.filter(code=code_projet).first()
                if projet_ligne is None:
                    erreurs.append({"ligne": numero, "message": f"Projet de code « {code_projet} » introuvable — ligne ignorée."})
                    continue

        niveau = str(valeur(ligne, "Niveau de rattachement (Aucun/Projet/Objectif général/Objectif spécifique/Activité)") or "")
        rattachement = {}
        if niveau.strip().upper() not in ("", "AUCUN"):
            if projet_ligne is None:
                erreurs.append({"ligne": numero, "message": "Un projet (code) est requis pour ce rattachement — ligne ignorée."})
                continue
            resolu = _resoudre_rattachement(projet_ligne, niveau, str(valeur(ligne, "Nom de l'élément rattaché") or ""))
            if resolu is None:
                avertissements.append({"ligne": numero, "message": "Élément de rattachement introuvable — indicateur créé sans rattachement."})
                resolu = {}
            rattachement = resolu

        valeur_reference = valeur(ligne, "Valeur de référence")
        indicateur = Indicateur.objects.create(
            libelle=str(libelle).strip(),
            unite=str(unite).strip(),
            valeur_reference=valeur_reference if valeur_reference not in (None, "") else None,
            valeur_cible=cible,
            frequence_collecte=frequence,
            **rattachement,
        )
        indicateurs_par_nom[indicateur.libelle.strip().lower()] = indicateur
        crees += 1

    return {
        "crees": crees,
        "erreurs": erreurs,
        "avertissements": avertissements,
        "_indicateurs_par_nom": indicateurs_par_nom,
    }
